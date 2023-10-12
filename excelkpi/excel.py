import openpyxl
import traceback
from openpyxl.worksheet.table import Table, TableColumn
from kpi import kpi_month
from datetime import datetime


class excel_kpi:
    currentWorkbook = None 
    ws = None

    # constructor
    def __init__(self, excel_path): 
        self.table_kpi_name_base = "KPIs"
        self.file_path = ""
        self.currentWorkbook = None
        self.ws = None
        self.tab = None
        self.filename = excel_path
        self.headers = ["DataDate","DateExe","SprintCompRate", "EscVelRate", "ReqInGherkin", "SprintReadiness", "SprintChurn", "TechDebtPaydown"]
        if len(excel_path) <= 0:
            raise print("file path can't be 0")
        try:
            self.currentWorkbook = openpyxl.load_workbook(filename=excel_path)
        except:
            print("couldn't open spreadsheet")
    
    def set_sheet(self, sheet_name):
        try:
            self.ws = self.currentWorkbook.get_sheet_by_name(sheet_name)
        except:
            print(f"set_sheet: {sheet_name} wasn't found, creating")
            self.ws = self.currentWorkbook.create_sheet(sheet_name)
        return True

    # Something in here excel isn't liking on file load after data is pushed
    def set_table(self, project, table_name):
        try:
            #todo - don't trust this  - search for it via
            self.tab = self.ws.tables[table_name]
        except:
            #style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            table = Table(ref='A1:H1',displayName=table_name,name=table_name)
            count = 1
            for header in self.headers:
                col = TableColumn(id=count, uniqueName=header, name=header)
                count = count+1
                table.tableColumns.append(col)
            self.tab = table
            self.ws.add_table(table)
            #self.save_file()
        return True

    def save_file(self):
        try:
            self.currentWorkbook.save(self.filename)
        except Exception:
            print(traceback.format_exc())
            return False
        return True

    #set the data to the table
    def set_data(self, kpi:kpi_month):

        #these all need to be unique - even across sheets
        if len(kpi.team_descriptor) > 0:
            sheet_name = f'{kpi.project}-{kpi.team_descriptor}'
            table_name = f'{kpi.project}-{self.table_kpi_name_base}-{kpi.team_descriptor}-table'
        else:
            sheet_name = f'{kpi.project}'
            table_name = f'{kpi.project}-{self.table_kpi_name_base}'
 
        self.set_sheet(sheet_name)
        data_date = str(kpi.year) + "." + str(kpi.month)
        curr_date = datetime.today().strftime('%Y.%m.%d')

        self.set_table(kpi.project, table_name)
        self.tab = self.ws.tables[table_name]
        self.ws.append([data_date, curr_date, kpi.Sprint_Completion_Rate, kpi.Escape_Velocity_Rate, kpi.Gherkin_Story_Rate, kpi.Sprint_Readiness_Ratio, kpi.Sprint_Churn, kpi.Tech_Debt_Paydown_Ratio])
        
        return
